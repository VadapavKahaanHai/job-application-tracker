import unittest
from datetime import date, timedelta
from io import BytesIO
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

import job_refresh as jobs


class JobRefreshTests(unittest.TestCase):
    today = date(2026, 10, 18)

    def result(self, **changes):
        result = dict(
            url="https://www.linkedin.com/jobs/view/12345",
            title="Werkstudent Data Analyst Hamburg",
            description="Python SQL; vor 2 Tagen; 0-2 years",
            markdown="Data preparation and reporting with Python and SQL.",
        )
        result.update(changes)
        return result

    def process(self, **changes):
        return jobs.process_results([self.result(**changes)], set(), self.today)

    def test_all_city_source_role_combinations(self):
        self.assertEqual(len(set(jobs.QUERIES)), 7 * 4 * 3)
        self.assertEqual(len(jobs.TARGET_CITIES), 7)
        for query, _ in jobs.QUERIES:
            self.assertLessEqual(len(query), 500)
            for skill in ("Python", "SQL", "PySpark", "PostgreSQL", "AWS", "Excel"):
                self.assertNotIn(skill, query)
            for keyword in ("Werkstudent", "Praktikum", "Internship"):
                self.assertIn(keyword, query)

    def test_sources_and_detail_pages(self):
        urls = {
            "https://de.linkedin.com/jobs/view/analyst-123": "LinkedIn",
            "https://hiring.cafe/viewjob/abc": "Hiring Cafe",
            "https://hiringcafe.com/viewjob/abc": "Hiring Cafe",
            "https://www.xing.com/jobs/hamburg-data-analyst-123": "XING",
            "https://www.stepstone.de/stellenangebote--Data-Analyst--123-inline.html": "StepStone",
        }
        for url, source in urls.items():
            with self.subTest(url=url):
                self.assertEqual(self.process(url=url)[0]["source"], source)
        for url in ("https://www.xing.com/jobs/t-data-analyst",
                    "https://stepstone.de/jobs/data-analyst",
                    "https://hiring.cafe/", "https://linkedin.com.evil.test/jobs/view/123"):
            self.assertEqual(self.process(url=url), [])

    def test_student_and_data_role_required(self):
        for title in ("Senior Data Analyst Hamburg", "Data Engineer Hamburg",
                      "Werkstudent Sales Hamburg", "International Data Analyst Hamburg",
                      "Intern Sales Hamburg"):
            self.assertEqual(self.process(title=title), [])
        for title in ("Werkstudentin Datenanalyse Hamburg", "Working Student Data Engineering Hamburg",
                      "Werkstudent Business Intelligence Hamburg", "Werkstudent:in Power BI Hamburg"):
            self.assertEqual(len(self.process(title=title)), 1)

    def test_internships_use_the_same_role_city_and_date_filters(self):
        for title in ("Intern Data Analyst Hamburg", "Internship Business Intelligence Hamburg",
                      "Praktikum Datenanalyse Hamburg", "Praktikant Data Engineering Hamburg",
                      "Praktikantin Power BI Hamburg", "Pflichtpraktikum Data Analytics Hamburg"):
            with self.subTest(title=title):
                self.assertEqual(len(self.process(title=title)), 1)
                self.assertEqual(self.process(title=title, description="vor 3 Tagen"), [])
                self.assertEqual(self.process(title=title.replace("Hamburg", "Berlin")), [])

    def test_cities_and_aliases(self):
        for alias, city in jobs.CITIES.items():
            self.assertEqual(self.process(title=f"Werkstudent Data Analyst {alias}")[0]["location"], city)
        for location in ("Berlin", "Remote", "Frankfurt an der Oder", "Frankfurt Oder"):
            self.assertEqual(self.process(title=f"Werkstudent Data Analyst {location}"), [])
        self.assertEqual(jobs.classify_location("", "https://example.test/k%C3%B6ln", ""), "Köln")
        self.assertEqual(self.process(title="Working Student Data Analyst", markdown="Location: Aachen")[0]["location"], "Aachen")

    def test_freshness_and_unknown_dates(self):
        for text, days in (("vor 2 Tagen", 2), ("vor 1 Woche", 7), ("3 days ago", 3),
                           ("vor 4 Stunden", 0), ("gestern", 1)):
            self.assertEqual(jobs.posted_date("", text, "", self.today), self.today - timedelta(days=days))
        for text in ("heute", "gestern", "vor 2 Tagen", "2 days ago", "71 hours ago"):
            self.assertEqual(len(self.process(description=text)), 1)
        for text in ("vor 3 Tagen", "3 days ago", "72 hours ago", "vor 72 Stunden",
                     "4320 minutes ago", "vor 4320 Minuten", "5 weeks ago", "1 month ago"):
            self.assertEqual(self.process(description=text), [])
        self.assertEqual(self.process(description="Python SQL"), [])

    def test_resume_fit_and_language_flag(self):
        self.assertEqual(self.process()[0]["fit"], "High")
        result = self.process(markdown="SQL Python. Sehr gute Deutschkenntnisse erforderlich.")[0]
        self.assertEqual(result["fit"], "Low")
        self.assertIn("B1", result["note"])
        self.assertEqual(jobs.fit_of("Python SQL; German B1"), "High")
        self.assertEqual(self.process()[0]["exp"], "0-2 yrs")

    def test_broader_roles_without_skills_are_kept(self):
        for role in ("Business Analytics", "Business Analyst", "Revenue Analytics",
                     "Marketing Analytics", "Datenanalytik", "BI Analyst", "BI Developer",
                     "Dashboard", "Datenvisualisierung", "Reporting Analyst",
                     "Analytics Engineering", "Data Platform", "Data Warehouse",
                     "Datenmanagement", "Datenintegration", "Datenqualität"):
            with self.subTest(role=role):
                self.assertTrue(any(role in query for query in jobs.ROLE_QUERIES))
                records = self.process(title=f"Praktikum {role} Hamburg", description="heute", markdown="")
                self.assertEqual(len(records), 1)
                self.assertEqual(records[0]["fit"], "Medium")

    def test_duplicates_and_snippet_fallback(self):
        seen = set()
        result = self.result(markdown="")
        first = jobs.process_results([result], seen, self.today)
        self.assertIn("snippet only", first[0]["note"])
        duplicate = dict(result, url=result["url"] + "?tracking=123#top")
        self.assertEqual(jobs.process_results([duplicate], seen, self.today), [])

    def test_search_requests_scraped_content_and_checks_errors(self):
        response = Mock()
        response.json.return_value = {"success": True, "data": {"web": [self.result()]}}
        with patch.object(jobs.requests, "post", return_value=response) as post:
            self.assertEqual(len(jobs.run_search("query", "Hamburg, Germany")), 1)
            body = post.call_args.kwargs["json"]
            self.assertEqual(body["country"], "DE")
            today = date.today()
            since = today - timedelta(days=2)
            self.assertEqual(body["tbs"], f"cdr:1,cd_min:{since:%m/%d/%Y},cd_max:{today:%m/%d/%Y}")
            self.assertEqual(body["scrapeOptions"]["formats"], [{"type": "markdown"}])
            response.json.return_value = {"success": False}
            with self.assertRaises(RuntimeError):
                jobs.run_search("query", "Hamburg, Germany")

    def test_excel_roundtrip_preserves_application_edits(self):
        wb = Workbook()
        ws = wb.active
        ws.append(jobs.HEADERS)
        records = self.process()
        records[0].update(status="Applied", applied=self.today, note="=literal note", fit="Low")
        jobs.write_records(ws, records)
        self.assertEqual(ws.cell(2, jobs.COL_NOTE).data_type, "s")
        self.assertIn("1 Target city", ws.cell(2, jobs.COL_PRI).value)
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        self.assertEqual(jobs.read_records(load_workbook(buffer).active), records)

    def test_all_failed_searches_do_not_save(self):
        wb = Workbook()
        wb.active.title = "Jobs"
        wb.active.append(jobs.HEADERS)
        with patch.object(jobs, "API_KEY", "test"), patch.object(jobs, "QUERIES", [("q", "city")]), \
                patch.object(jobs, "load_workbook", return_value=wb), \
                patch.object(jobs, "run_search", side_effect=RuntimeError("offline")), \
                patch.object(wb, "save") as save:
            with self.assertRaises(SystemExit):
                jobs.main()
            save.assert_not_called()


if __name__ == "__main__":
    unittest.main()
